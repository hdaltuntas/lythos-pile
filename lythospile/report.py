"""
Calculation report for Lythos Pile.

The report is assembled as HTML (the inputs, the base and shaft resistance of
every method, the pile's weight, the capacity, the group, the settlement, the
checks and the required length, the figures, the warnings and the method
notes; the rock socket if one was analysed; the study if one was run) and
exported as
  * PDF   — through reportlab (see `lythospile.pdf`)
  * HTML  — a single self-contained file (figures embedded as base64)
  * DOCX  — optional, needs `python-docx`

There is one assembly, `build_html()`, so all three formats say the same thing.
"""

from __future__ import annotations

import base64
import datetime
import html
import io
import math
from typing import Any, Dict, List, Optional

from matplotlib.figure import Figure

from . import study_plots
from .config import (
    APP_NAME,
    APP_VERSION,
    CLAY_METHODS,
    SOCKET_BASE_METHODS,
    SOCKET_SIDE_METHODS,
    TIP_METHODS,
)
from .i18n import TRANSLATIONS, warning_text
from .plotting import PLOT_KEYS, SOCKET_PLOT_KEYS, Plotter, SocketPlotter
from .study import OUTPUTS
from .summary import method_label

TEXTS = {
    "en": {
        "title": f"{APP_NAME} — Calculation Report",
        "date": "Date", "analyst": "Analyst", "software": "Software",
        "sec_inputs": "Input data", "sec_pile": "Pile, group and load",
        "sec_soil": "Soil profile", "sec_opts": "Methods, options and criteria",
        "sec_single": "Capacity of a single pile", "sec_base": "Base resistance",
        "sec_shaft_layers": "Shaft friction by layer",
        "sec_shaft_methods": "Shaft friction by method",
        "sec_capacity": "Ultimate and allowable capacity",
        "sec_matrix": "Net ultimate capacity, shaft method × base method",
        "sec_group": "Pile group", "sec_eff": "Group efficiency", "sec_block": "Block failure",
        "sec_group_cap": "Capacity of the group",
        "sec_settlement": "Settlement", "sec_single_s": "Single pile (Vesić)",
        "sec_group_s": "Group", "sec_raft": "Equivalent raft, slice by slice",
        "sec_checks": "Checks and the required length",
        "sec_figs": "Figures", "sec_warn": "Warnings", "sec_notes": "Method notes",
        "sec_socket": "Rock-socketed pile", "sec_socket_in": "Socket, rock and concrete",
        "sec_socket_side": "Unit side shear by correlation",
        "sec_socket_base": "Unit base resistance",
        "sec_socket_design": "Design, socket length and check",
        "sec_socket_s": "Elastic settlement", "sec_socket_figs": "Figures",
        "sec_socket_warn": "Warnings", "sec_socket_notes": "Method notes",
        "parameter": "Parameter", "value": "Value", "unit": "Unit",
        "shape": "Shape", "D": "Diameter / side D", "L": "Length L",
        "top": "Depth of the pile head", "tip": "Depth of the tip",
        "installation": "Installation", "gamma_p": "Unit weight of the pile γp",
        "Ep": "Modulus of the pile Ep", "A": "Base area Ab", "p": "Perimeter",
        "Q": "Load on the group Q", "n": "Number of piles", "Qp": "Load per pile",
        "sx": "Spacing along B", "sy": "Spacing along L",
        "zw": "Water table depth", "gw": "Unit weight of water γw",
        "yes": "yes", "no": "no",
        "layer": "Layer", "type": "Type", "top_l": "Top", "bottom": "Bottom",
        "clay_method": "Shaft friction in clay", "tip_method": "Base resistance",
        "K": "K / K0 in sand", "delta": "δ / φ'", "zc": "Critical depth zc",
        "weight": "Pile weight subtracted", "buoyant": "Buoyant below the water table",
        "efficiency": "Group efficiency", "block": "Block failure checked",
        "gs": "Group settlement for the check", "FS": "Factor of safety",
        "s_allow": "Allowable settlement",
        "method": "Method", "factor": "Factor", "qb": "qb (kPa)", "Qb": "Qb (kN)",
        "depth": "Depth (m)", "Qs": "Qs (kN)", "sv": "σ'v0 at the tip",
        "Qs_l": "Shaft friction Qs", "Qb_l": "Base resistance Qb", "Qu": "Qult = Qs + Qb",
        "W": "Pile weight W", "Qn": "Qult,net = Qult − W", "Qa": "Qall = Qult,net / FS",
        "eta": "η", "Bg": "Width of the block Bg", "Lg": "Length of the block Lg",
        "block_Qs": "Shaft of the block", "block_Qb": "Base of the block",
        "block_Q": "Capacity of the block", "Q_eff": "η · n · Qult",
        "Qg": "Qg,ult", "Qgn": "Qg,ult − n·W", "Qga": "Qg,all", "governs": "Governs",
        "s1": "Shortening of the pile s1", "s2": "Base load s2", "s3": "Shaft load s3",
        "s": "Settlement", "z": "z (m)", "h": "h (m)", "s0": "σ'v0 (kPa)",
        "ds": "Δσ (kPa)", "ds_mm": "s (mm)", "kind": "Model",
        "raft_z": "Depth of the raft", "raft_q": "Pressure on the raft",
        "raft_c": "Consolidation", "raft_e": "Elastic compression",
        "raft_sh": "Shortening of the piles above the raft", "raft_total": "Total",
        "check": "Check", "actual": "Actual", "allowable": "Allowable", "status": "Status",
        "c_pile": "Single pile (FS)", "c_group": "Group (FS)", "c_settle": "Settlement (mm)",
        "ok": "OK", "notok": "NOT OK", "na": "n/a",
        "socket_D": "Socket diameter D", "socket_Ls": "Socket length checked Ls",
        "socket_top": "Depth of the pile head", "rock": "Depth of the rock surface",
        "overburden": "Length through the overburden", "socket_Q": "Load Q",
        "qu": "Intact rock strength qu", "qu_side": "qu for the side shear (≤ f'c)",
        "RQD": "RQD", "Ei": "Intact modulus Ei", "Em": "Rock mass modulus Em",
        "ratio": "Em / Ei", "alpha_E": "Joint factor αE", "GSI": "GSI", "mi": "mi",
        "mb": "mb", "s_hb": "s", "fc": "Concrete strength f'c", "Ec": "Concrete modulus Ec",
        "correlation": "Correlation", "formula": "fs (MPa)", "fs": "fs (kPa)",
        "Ls_req": "Ls needed (m)", "Q_all": "Qall at Ls (kN)", "range": "In range",
        "base_formula": "qb (MPa)", "qb_mpa": "qb (MPa)",
        "fs_design": "Design side shear", "qb_design": "Design base resistance",
        "FS_side": "Factor of safety on the side", "FS_base": "Factor of safety on the base",
        "Ls_need": "Socket length needed", "Ls_min": "Minimum socket length",
        "Ls_design": "Design socket length",
        "rw": "Randolph & Wroth, side and base", "rw_side": "Randolph & Wroth, side only",
        "vesic": "Vesić", "shortening": "Shortening through the overburden",
        "base_share": "Share carried by the base",
        "sec_study": "Parametric / reliability study", "sec_study_vars": "Variables",
        "sec_study_stats": "Statistics of the outputs",
        "sec_study_rel": "Probability of failure",
        "sec_study_sens": "Sensitivities (Spearman ρ)", "sec_study_figs": "Figures",
        "study_method": "Sampling method", "study_n": "Samples",
        "study_ok": "Successful analyses", "variable": "Variable", "mode": "Mode",
        "range_mode": "Range", "dist": "Distribution", "mean": "Mean", "cov": "CoV",
        "min": "Min", "max": "Max", "output": "Output", "std": "Std",
        "criterion": "Criterion", "n_fail": "Failed", "pf": "P", "pf_ci": "95 % CI",
        "beta_idx": "β",
        "notes": [
            "The ultimate capacity of a pile is Qult = Qs + Qb: the shaft friction "
            "integrated down the shaft, Qs = Σ fs·p·Δz, and the base resistance Qb = qb·Ab. "
            "The pile's weight W = γp·Ab·L, buoyant (γp − γw) below the water table, is "
            "taken off: Qult,net = Qult − W, and Qall = Qult,net / FS.",
            "Stresses are effective: σ'v0 = σv0 − u0, u0 hydrostatic below the water "
            "table. In sand the shaft friction is fs = K·σ'v·tan δ with K = (K/K0)·(1 − sin φ'), "
            "and σ'v is held constant below the critical depth zc when asked (Meyerhof; "
            "Vesić).",
            "In clay: API RP 2A α = 0.5·ψ^-0.5 (ψ ≤ 1) or 0.5·ψ^-0.25, ψ = cu/σ'v; Kulhawy & "
            "Phoon α = 0.21 + 0.26·pa/cu; Sladen α = C·(σ'v/cu)^0.45; each α ≤ 1 and "
            "fs = α·cu. The β method fs = (1 − sin φ')·tan φ'·√OCR·σ'v; the λ method "
            "fs = λ·(σ'v + 2cu), λ from the pile's penetration (Vijayvergiya & Focht).",
            "The base in sand: Meyerhof qb = σ'v·Nq* ≤ 0.5·pa·Nq*·tan φ'; Vesić "
            "qb = σ'v·Nq*(Irr), Irr the rigidity index reduced for volume change; Janbu "
            "qb = σ'v·(tan φ' + √(1 + tan² φ'))²·e^(2η'·tan φ'). In clay qb = Nc*·cu with "
            "Nc* = 9, Vesić's 4/3·(ln Ir + 1) + π/2 + 1, or Janbu's 2 + 2η'; the overburden "
            "term is left out, as it is balanced by the pile's own weight.",
            "The group's capacity is the smaller of η·n·Qult and the capacity of the block "
            "the group occupies (soil-on-soil shaft, the base resistance over the footprint; "
            "Skempton's Nc in clay); the piles' weight n·W is then taken off.",
            "The settlement of one pile is Vesić's s1 + s2 + s3, the working load shared "
            "between shaft and base in the proportion of their ultimate resistances. The "
            "group settles as an equivalent raft at 2/3 of the pile length, the load spread "
            "2 : 1 below it (clays by Cc, Cr, e0 and OCR, the others by their constrained "
            "modulus), plus the piles' shortening above it; by Vesić's s·√(Bg/D); or by "
            "Meyerhof's SPT rule in sand.",
            "The required length is the shortest pile, in steps of the length search, "
            "that passes the single-pile and the group check with every other input held "
            "as it is.",
        ],
        "socket_notes": [
            "The side shear is fs·π·D·Ls and the base resistance qb·π·D²/4. The allowable "
            "capacity is Qs/FSside + Qb/FSbase − W; the socket length is the shortest that "
            "carries the load, found by bisection, and at least the minimum length.",
            "qu in the side shear correlations is the weaker of the rock and the concrete. "
            "The linear rules (Reynolds & Kaderabek, Gupton & Logan, Toh et al.) were "
            "fitted to weak rock and are left out of the statistics above the weak rock "
            "limit.",
            "O'Neill & Reese's αE reduces the side shear of jointed rock through Em/Ei; "
            "the rock mass modulus comes from RQD (Gardner 1987), from GSI (Hoek & "
            "Diederichs 2006), or is entered.",
            "The settlement is Randolph & Wroth's closed form for a compressible pile in an "
            "elastic medium, with the socket's rock along the shaft and below the base, plus "
            "the elastic shortening of the pile through the overburden, whose friction is "
            "ignored.",
        ],
    },
    "tr": {
        "title": f"{APP_NAME} — Hesap Raporu",
        "date": "Tarih", "analyst": "Hazırlayan", "software": "Yazılım",
        "sec_inputs": "Girdi verileri", "sec_pile": "Kazık, grup ve yük",
        "sec_soil": "Zemin profili", "sec_opts": "Yöntemler, seçenekler ve ölçütler",
        "sec_single": "Tek kazığın taşıma gücü", "sec_base": "Uç direnci",
        "sec_shaft_layers": "Tabakaya göre çevre sürtünmesi",
        "sec_shaft_methods": "Yönteme göre çevre sürtünmesi",
        "sec_capacity": "Nihai ve izin verilebilir taşıma gücü",
        "sec_matrix": "Net nihai taşıma gücü, sürtünme yöntemi × uç yöntemi",
        "sec_group": "Kazık grubu", "sec_eff": "Grup verimi", "sec_block": "Blok göçmesi",
        "sec_group_cap": "Grubun taşıma gücü",
        "sec_settlement": "Oturma", "sec_single_s": "Tek kazık (Vesić)",
        "sec_group_s": "Grup", "sec_raft": "Eşdeğer radye, dilim dilim",
        "sec_checks": "Kontroller ve gerekli boy",
        "sec_figs": "Şekiller", "sec_warn": "Uyarılar", "sec_notes": "Yöntem notları",
        "sec_socket": "Kayaya soketli kazık", "sec_socket_in": "Soket, kaya ve beton",
        "sec_socket_side": "Bağıntılara göre birim yanal sürtünme",
        "sec_socket_base": "Birim uç direnci",
        "sec_socket_design": "Tasarım, soket boyu ve kontrol",
        "sec_socket_s": "Elastik oturma", "sec_socket_figs": "Şekiller",
        "sec_socket_warn": "Uyarılar", "sec_socket_notes": "Yöntem notları",
        "parameter": "Parametre", "value": "Değer", "unit": "Birim",
        "shape": "Kesit", "D": "Çap / kenar D", "L": "Boy L",
        "top": "Kazık başı derinliği", "tip": "Uç derinliği",
        "installation": "İmalat", "gamma_p": "Kazık malzemesinin birim hacim ağırlığı γp",
        "Ep": "Kazık malzemesinin elastisite modülü Ep", "A": "Uç alanı Ab", "p": "Çevre",
        "Q": "Gruba gelen yük Q", "n": "Kazık sayısı", "Qp": "Kazık başına yük",
        "sx": "B yönünde aralık", "sy": "L yönünde aralık",
        "zw": "Su tablası derinliği", "gw": "Suyun birim hacim ağırlığı γw",
        "yes": "evet", "no": "hayır",
        "layer": "Tabaka", "type": "Tür", "top_l": "Üst", "bottom": "Alt",
        "clay_method": "Kilde çevre sürtünmesi", "tip_method": "Uç direnci",
        "K": "Kumda K / K0", "delta": "δ / φ'", "zc": "Kritik derinlik zc",
        "weight": "Kazık ağırlığı düşüldü", "buoyant": "Su altında batık ağırlık",
        "efficiency": "Grup verimi", "block": "Blok göçmesi kontrol edildi",
        "gs": "Kontrolde kullanılan grup oturması", "FS": "Güvenlik sayısı",
        "s_allow": "İzin verilen oturma",
        "method": "Yöntem", "factor": "Katsayı", "qb": "qb (kPa)", "Qb": "Qb (kN)",
        "depth": "Derinlik (m)", "Qs": "Qs (kN)", "sv": "Uçta σ'v0",
        "Qs_l": "Çevre sürtünmesi Qs", "Qb_l": "Uç direnci Qb", "Qu": "Qult = Qs + Qb",
        "W": "Kazık ağırlığı W", "Qn": "Qult,net = Qult − W", "Qa": "Qall = Qult,net / GS",
        "eta": "η", "Bg": "Blok genişliği Bg", "Lg": "Blok boyu Lg",
        "block_Qs": "Blok çevresi", "block_Qb": "Blok tabanı",
        "block_Q": "Bloğun taşıma gücü", "Q_eff": "η · n · Qult",
        "Qg": "Qg,ult", "Qgn": "Qg,ult − n·W", "Qga": "Qg,all", "governs": "Belirleyici",
        "s1": "Kazığın kısalması s1", "s2": "Uç yükünden s2", "s3": "Çevre yükünden s3",
        "s": "Oturma", "z": "z (m)", "h": "h (m)", "s0": "σ'v0 (kPa)",
        "ds": "Δσ (kPa)", "ds_mm": "s (mm)", "kind": "Model",
        "raft_z": "Radye derinliği", "raft_q": "Radyeye gelen basınç",
        "raft_c": "Konsolidasyon", "raft_e": "Elastik sıkışma",
        "raft_sh": "Radye üstündeki kazık kısalması", "raft_total": "Toplam",
        "check": "Kontrol", "actual": "Oluşan", "allowable": "İzin verilen", "status": "Durum",
        "c_pile": "Tek kazık (GS)", "c_group": "Grup (GS)", "c_settle": "Oturma (mm)",
        "ok": "UYGUN", "notok": "UYGUN DEĞİL", "na": "—",
        "socket_D": "Soket çapı D", "socket_Ls": "Kontrol edilen soket boyu Ls",
        "socket_top": "Kazık başı derinliği", "rock": "Kaya yüzeyinin derinliği",
        "overburden": "Örtü tabakasındaki boy", "socket_Q": "Yük Q",
        "qu": "Sağlam kaya dayanımı qu", "qu_side": "Yanal sürtünme için qu (≤ f'c)",
        "RQD": "RQD", "Ei": "Sağlam kaya modülü Ei", "Em": "Kaya kütlesi modülü Em",
        "ratio": "Em / Ei", "alpha_E": "Süreksizlik katsayısı αE", "GSI": "GSI", "mi": "mi",
        "mb": "mb", "s_hb": "s", "fc": "Beton dayanımı f'c", "Ec": "Beton modülü Ec",
        "correlation": "Bağıntı", "formula": "fs (MPa)", "fs": "fs (kPa)",
        "Ls_req": "Gereken Ls (m)", "Q_all": "Ls'de Qall (kN)", "range": "Geçerli",
        "base_formula": "qb (MPa)", "qb_mpa": "qb (MPa)",
        "fs_design": "Tasarım yanal sürtünmesi", "qb_design": "Tasarım uç direnci",
        "FS_side": "Yanal sürtünme güvenlik sayısı", "FS_base": "Uç direnci güvenlik sayısı",
        "Ls_need": "Gereken soket boyu", "Ls_min": "En küçük soket boyu",
        "Ls_design": "Tasarım soket boyu",
        "rw": "Randolph & Wroth, çevre ve uç", "rw_side": "Randolph & Wroth, yalnız çevre",
        "vesic": "Vesić", "shortening": "Örtü tabakasındaki kısalma",
        "base_share": "Ucun taşıdığı pay",
        "sec_study": "Parametrik / güvenilirlik çalışması", "sec_study_vars": "Değişkenler",
        "sec_study_stats": "Çıktıların istatistikleri", "sec_study_rel": "Göçme olasılığı",
        "sec_study_sens": "Duyarlılıklar (Spearman ρ)", "sec_study_figs": "Şekiller",
        "study_method": "Örnekleme yöntemi", "study_n": "Örnek sayısı",
        "study_ok": "Başarılı analiz", "variable": "Değişken", "mode": "Mod",
        "range_mode": "Aralık", "dist": "Dağılım", "mean": "Ortalama", "cov": "CoV",
        "min": "Min", "max": "Maks", "output": "Çıktı", "std": "Std",
        "criterion": "Ölçüt", "n_fail": "Göçen", "pf": "P", "pf_ci": "%95 GA",
        "beta_idx": "β",
        "notes": [
            "Kazığın nihai taşıma gücü Qult = Qs + Qb'dir: gövde boyunca toplanan çevre "
            "sürtünmesi Qs = Σ fs·p·Δz ve uç direnci Qb = qb·Ab. Kazık ağırlığı W = γp·Ab·L, "
            "su tablası altında batık (γp − γw) olarak düşülür: Qult,net = Qult − W ve "
            "Qall = Qult,net / GS.",
            "Gerilmeler efektiftir: σ'v0 = σv0 − u0, su tablası altında u0 hidrostatiktir. "
            "Kumda çevre sürtünmesi fs = K·σ'v·tan δ, K = (K/K0)·(1 − sin φ'); istenirse σ'v "
            "kritik derinlik zc'nin altında sabit tutulur (Meyerhof; Vesić).",
            "Kilde: API RP 2A α = 0.5·ψ^-0.5 (ψ ≤ 1) ya da 0.5·ψ^-0.25, ψ = cu/σ'v; Kulhawy & "
            "Phoon α = 0.21 + 0.26·pa/cu; Sladen α = C·(σ'v/cu)^0.45; her α ≤ 1 ve fs = α·cu. "
            "β yöntemi fs = (1 − sin φ')·tan φ'·√OCR·σ'v; λ yöntemi fs = λ·(σ'v + 2cu), λ "
            "kazığın gömülme derinliğinden (Vijayvergiya & Focht).",
            "Kumda uç: Meyerhof qb = σ'v·Nq* ≤ 0.5·pa·Nq*·tan φ'; Vesić qb = σ'v·Nq*(Irr), "
            "Irr hacim değişimiyle azaltılmış rijitlik indisi; Janbu "
            "qb = σ'v·(tan φ' + √(1 + tan² φ'))²·e^(2η'·tan φ'). Kilde qb = Nc*·cu; Nc* = 9, "
            "Vesić'in 4/3·(ln Ir + 1) + π/2 + 1'i ya da Janbu'nun 2 + 2η''sü; örtü yükü terimi, "
            "kazığın kendi ağırlığıyla dengelendiği için alınmaz.",
            "Grubun taşıma gücü η·n·Qult ile grubun oluşturduğu bloğun taşıma gücünün "
            "küçüğüdür (zemin-zemin çevre sürtünmesi, taban alanı boyunca uç direnci; kilde "
            "Skempton Nc); ardından kazıkların ağırlığı n·W düşülür.",
            "Tek kazığın oturması Vesić'in s1 + s2 + s3'üdür; çalışma yükü çevre ve uç "
            "arasında nihai dirençleri oranında paylaştırılır. Grup, kazık boyunun 2/3'ündeki "
            "eşdeğer radye olarak oturur; yük altında 2 : 1 yayılır (killer Cc, Cr, e0 ve OCR "
            "ile, diğerleri sınırlandırılmış modülleriyle), radye üstündeki kazık kısalması "
            "eklenir; ya da Vesić'in s·√(Bg/D)'si; ya da kumda Meyerhof'un SPT kuralı.",
            "Gerekli boy, diğer bütün girdiler sabit tutularak tek kazık ve grup "
            "kontrollerini sağlayan, boy aramasının adımlarındaki en kısa kazıktır.",
        ],
        "socket_notes": [
            "Yanal sürtünme fs·π·D·Ls, uç direnci qb·π·D²/4'tür. İzin verilebilir taşıma gücü "
            "Qs/GSyanal + Qb/GSuç − W; soket boyu yükü taşıyan en kısa boydur, ikiye bölme ile "
            "bulunur ve en küçük boydan az olamaz.",
            "Yanal sürtünme bağıntılarındaki qu, kaya ile betonun zayıf olanıdır. Doğrusal "
            "kurallar (Reynolds & Kaderabek, Gupton & Logan, Toh vd.) zayıf kayaya "
            "uydurulmuştur; zayıf kaya sınırının üstünde istatistiklere katılmaz.",
            "O'Neill & Reese'in αE katsayısı, eklemli kayanın yanal sürtünmesini Em/Ei "
            "üzerinden azaltır; kaya kütlesi modülü RQD'den (Gardner 1987), GSI'dan (Hoek & "
            "Diederichs 2006) ya da doğrudan girilir.",
            "Oturma, elastik ortamdaki sıkışabilir kazık için Randolph & Wroth'un kapalı "
            "çözümüdür; gövde boyunca ve uç altında soketin kayası alınır; sürtünmesi ihmal "
            "edilen örtü tabakasındaki kazığın elastik kısalması eklenir.",
        ],
    },
}

#: The interface's palette: warm paper, ink, terracotta; serif headings
_CSS = """
body { font-family: system-ui, -apple-system, 'Segoe UI', 'Helvetica Neue', Arial, sans-serif;
       font-size: 10pt; color: #141413; }
h1, h2, h3 { font-family: 'Tiempos Text', 'Source Serif 4', 'Iowan Old Style', Palatino,
             Georgia, 'DejaVu Serif', serif; font-weight: 500; }
h1 { font-size: 20pt; color: #141413; margin-bottom: 2px; }
h2 { font-size: 14pt; color: #C6613F; border-bottom: 1px solid #E3E0D5; padding-bottom: 2px;
     margin-top: 20px; }
h3 { font-size: 11.5pt; color: #141413; margin-top: 12px; }
table { border-collapse: collapse; margin: 4px 0 8px 0; }
th { background: #F0EEE6; text-align: left; padding: 3px 6px; border: 1px solid #E3E0D5;
     font-size: 9pt; }
td { padding: 3px 6px; border: 1px solid #E3E0D5; font-size: 9pt; }
.ok { color: #3F7F4F; font-weight: bold; } .bad { color: #B0413E; font-weight: bold; }
.meta { color: #73726C; } .note { color: #73726C; font-size: 9pt; }
"""


def _esc(x) -> str:
    return html.escape(str(x))


def _f(x, nd=2) -> str:
    if x is None or (isinstance(x, float) and not math.isfinite(x)):
        return "—"
    return f"{x:,.{nd}f}"


def _status(T, s: str) -> str:
    if s == "N/A":
        return T["na"]
    return (f'<span class="ok">{T["ok"]}</span>' if s == "OK"
            else f'<span class="bad">{T["notok"]}</span>')


def _table(headers: List[str], rows: List[List[Any]], widths: Optional[List[int]] = None) -> str:
    out = ["<table width='100%'>"]
    if widths:
        cells = "".join(f"<th width='{w}%'>{_esc(h)}</th>" for h, w in zip(headers, widths))
    else:
        cells = "".join(f"<th>{_esc(h)}</th>" for h in headers)
    out.append("<tr>" + cells + "</tr>")
    for r in rows:
        out.append("<tr>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>")
    out.append("</table>")
    return "\n".join(out)


def _kv_table(T, rows: List[List[Any]]) -> str:
    return _table([T["parameter"], T["value"], T["unit"]], rows, [50, 36, 14])


class _Numbers:
    """Section numbers that follow whatever the report happens to contain."""

    def __init__(self):
        self.major = 0
        self.minor = 0

    def h2(self, text: str, page_break: bool = False) -> str:
        self.major += 1
        self.minor = 0
        style = " style='page-break-before:always'" if page_break else ""
        return f"<h2{style}>{self.major}. {_esc(text)}</h2>"

    def h3(self, text: str) -> str:
        self.minor += 1
        return f"<h3>{self.major}.{self.minor} {_esc(text)}</h3>"


# ----------------------------------------------------------------------
def _png(fig: Figure, dpi: int) -> bytes:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=dpi, bbox_inches="tight")
    return buf.getvalue()


def render_figures(analysis, lang: str, keys=PLOT_KEYS, dpi: int = 130) -> Dict[str, bytes]:
    """The report figures of the pile analysis as PNG bytes (off-screen, white, untitled)."""
    if analysis is None:
        return {}
    plotter = Plotter(analysis, lang, "paper", titles=False)
    out = {}
    for key in keys:
        fig = Figure(figsize=(10, 7), dpi=dpi)
        plotter.draw(key, fig)
        out[key] = _png(fig, dpi)
    return out


def render_socket_figures(socket, lang: str, dpi: int = 130) -> Dict[str, bytes]:
    if socket is None:
        return {}
    plotter = SocketPlotter(socket, lang, "paper", titles=False)
    out = {}
    for key in SOCKET_PLOT_KEYS:
        fig = Figure(figsize=(10, 7), dpi=dpi)
        plotter.draw(key, fig)
        out[key] = _png(fig, dpi)
    return out


def render_study_figures(study, lang: str, dpi: int = 110) -> Dict[str, bytes]:
    """Study figures keyed 'study_<view>' (only the views that apply)."""
    if study is None or not study.rows:
        return {}
    L = TRANSLATIONS[lang]
    views = ["oat", "hist"] if study.method == "oat" else ["hist", "scatter", "tornado"]
    draw = {"oat": study_plots.plot_oat, "hist": study_plots.plot_hist,
            "scatter": study_plots.plot_scatter, "tornado": study_plots.plot_tornado}
    out = {}
    for view in views:
        fig = Figure(figsize=(10, 7), dpi=dpi)
        draw[view](fig, study, L, "FS", theme="paper")
        out[f"study_{view}"] = _png(fig, dpi)
    return out


# ----------------------------------------------------------------------
def _pile_section(T, L, a, figures, img_src, lang, N: _Numbers) -> List[str]:
    res = a.results
    parts = [N.h2(T["sec_inputs"]), N.h3(T["sec_pile"])]
    rows = [[T["shape"], L["shape_" + a.shape], ""], [T["D"], _f(a.D), "m"],
            [T["L"], _f(a.L), "m"], [T["top"], _f(a.top), "m"], [T["tip"], _f(res["z_tip"]), "m"],
            [T["installation"], L["installation_" + a.installation], ""],
            [T["gamma_p"], _f(a.gamma_p, 1), "kN/m³"], [T["Ep"], _f(a.Ep, 0), "MPa"],
            [T["A"], _f(res["A"], 4), "m²"], [T["p"], _f(res["perimeter"], 3), "m"],
            [T["Q"], _f(a.Q, 0), "kN"], [T["n"], f"{a.nx} × {a.ny} = {a.n}", ""],
            [T["Qp"], _f(res["Q_pile"], 0), "kN"]]
    if a.nx > 1:
        rows.append([T["sx"], _f(a.sx), "m"])
    if a.ny > 1:
        rows.append([T["sy"], _f(a.sy), "m"])
    rows += [[T["zw"], _f(a.profile.zw), "m"], [T["gw"], _f(a.profile.gw), "kN/m³"]]
    parts.append(_kv_table(T, rows))

    parts.append(N.h3(T["sec_soil"]))
    rows = [[_esc(r["name"]), L["behaviour_" + r["behaviour"]], _f(r["top"]), _f(r["bottom"]),
             _f(r["gamma"], 1), _f(r["gamma_sat"], 1), _f(r["phi"], 1),
             _f(r["cu"], 0) if r["behaviour"] == "cohesive" else "—",
             _f(r["OCR"], 1) if r["behaviour"] == "cohesive" else "—",
             _f(r["N60"], 0) if r["N60"] > 0 else "—", _f(r["E"], 0),
             _f(r["sigma_eff"], 1)] for r in res["layers"]]
    parts.append(_table([T["layer"], T["type"], T["top_l"], T["bottom"], L["col_gamma"],
                         L["col_gamma_sat"], L["col_phi"], L["col_cu"], L["col_OCR"],
                         L["col_N60"], L["col_E"], "σ'v0"], rows,
                        [15, 9, 6, 6, 8, 8, 7, 7, 6, 7, 8, 9]))

    parts.append(N.h3(T["sec_opts"]))
    rows = [[T["clay_method"], method_label(L, a.clay_method), ""],
            [T["tip_method"], method_label(L, a.tip_method), ""],
            [T["K"], _f(a.K_ratio), ""], [T["delta"], _f(a.delta_ratio), ""],
            [T["zc"], _f(res["zc"]) if math.isfinite(res["zc"]) else T["no"],
             "m" if math.isfinite(res["zc"]) else ""],
            [T["weight"], T["yes"] if a.subtract_weight else T["no"], ""],
            [T["buoyant"], T["yes"] if a.buoyant_weight else T["no"], ""],
            [T["efficiency"], L["eff_" + a.efficiency_method], ""],
            [T["block"], T["yes"] if a.use_block else T["no"], ""],
            [T["gs"], L["gs_" + a.group_settlement], ""],
            [T["FS"], _f(a.FS), ""], [T["s_allow"], _f(a.s_allow, 0), "mm"]]
    parts.append(_kv_table(T, rows))

    # ---------------------------------------------------------- single pile
    parts.append(N.h2(T["sec_single"]))
    base = res["base"]
    parts.append(N.h3(T["sec_base"]))
    parts.append(f"<p>{_esc(base['layer'])} · z = {_f(base['z'])} m · {T['sv']} = "
                 f"{_f(base['sv'], 1)} kPa</p>")
    rows = []
    for tip in TIP_METHODS:
        m = base["methods"][tip]
        name, value = ("Nc*", m["Nc"]) if "Nc" in m else ("Nq*", m["Nq"])
        mark = " *" if tip == a.tip_method else ""
        rows.append([_esc(method_label(L, tip) + mark), f"{name} = {_f(value, 1)}",
                     _f(m["qb"], 0), _f(m["Qb"], 0)])
    if math.isfinite(base["Qb_spt"]):
        rows.append([_esc(method_label(L, "spt")), "—", _f(base["qb_spt"], 0),
                     _f(base["Qb_spt"], 0)])
    parts.append(_table([T["method"], T["factor"], T["qb"], T["Qb"]], rows, [40, 20, 20, 20]))

    parts.append(N.h3(T["sec_shaft_layers"]))
    rows = []
    for layer, lo, hi in a.profile.segments(a.top, res["z_tip"]):
        part = sum(row[a.clay_method] * res["perimeter"] * row["h"]
                   for row in res["shaft"]["slices"] if lo - 1e-9 <= row["z"] <= hi + 1e-9)
        rows.append([_esc(layer["name"]), L["behaviour_" + layer["behaviour"]],
                     f"{lo:.2f} – {hi:.2f}", _f(part, 0)])
    parts.append(_table([T["layer"], T["type"], T["depth"], T["Qs"]], rows, [40, 20, 20, 20]))
    if res["shaft"]["cohesive"] or math.isfinite(res["shaft"]["Qs_spt"]):
        parts.append(N.h3(T["sec_shaft_methods"]))
        rows = [[_esc(method_label(L, m) + (" *" if m == a.clay_method else "")),
                 _f(res["shaft"]["Qs"][m], 0)]
                for m in (CLAY_METHODS if res["shaft"]["cohesive"] else [])]
        if math.isfinite(res["shaft"]["Qs_spt"]):
            rows.append([_esc(method_label(L, "spt")), _f(res["shaft"]["Qs_spt"], 0)])
        parts.append(_table([T["method"], T["Qs"]], rows, [60, 40]))

    parts.append(N.h3(T["sec_capacity"]))
    parts.append(_kv_table(T, [
        [T["Qs_l"], _f(res["Qs"], 0), "kN"], [T["Qb_l"], _f(res["Qb"], 0), "kN"],
        [T["Qu"], _f(res["Q_ult"], 0), "kN"],
        [T["W"], _f(res["W"], 0) + ("" if a.subtract_weight else f" ({T['no']})"), "kN"],
        [f"<b>{T['Qn']}</b>", f"<b>{_f(res['Q_ult_net'], 0)}</b>", "kN"],
        [f"<b>{T['Qa']}</b>", f"<b>{_f(res['Q_all'], 0)}</b>", "kN"],
    ]))
    parts.append(N.h3(T["sec_matrix"]))
    clays = CLAY_METHODS if res["shaft"]["cohesive"] else [a.clay_method]
    rows = [[_esc(method_label(L, c) if res["shaft"]["cohesive"] else L["method_granular"])]
            + [_f(res["combos"][f"{c}|{t}"]["Q_ult_net"], 0) for t in TIP_METHODS]
            for c in clays]
    parts.append(_table([T["method"]] + [method_label(L, t) for t in TIP_METHODS], rows,
                        [34, 22, 22, 22]))

    # ---------------------------------------------------------- group
    group = res["group"]
    parts.append(N.h2(T["sec_group"]))
    parts.append(N.h3(T["sec_eff"]))
    rows = [[_esc(L["eff_" + k] + (" *" if k == a.efficiency_method else "")), _f(v, 3)]
            for k, v in group["eta"].items()]
    parts.append(_table([T["method"], T["eta"]], rows, [70, 30]))
    if group["block"] is not None:
        b = group["block"]
        parts.append(N.h3(T["sec_block"]))
        parts.append(_kv_table(T, [[T["Bg"], _f(b["Bg"]), "m"], [T["Lg"], _f(b["Lg"]), "m"],
                                   [T["block_Qs"], _f(b["Qs"], 0), "kN"],
                                   [T["block_Qb"], _f(b["Qb"], 0), "kN"],
                                   [T["block_Q"], _f(b["Q_ult"], 0), "kN"]]))
    parts.append(N.h3(T["sec_group_cap"]))
    parts.append(_kv_table(T, [
        [T["Q_eff"], _f(group["Q_eff"], 0), "kN"],
        [T["Qg"], _f(group["Q_ult"], 0), "kN"],
        [T["governs"], L["gov_" + group["governing"]], ""],
        [T["Qgn"], _f(group["Q_ult_net"], 0), "kN"],
        [f"<b>{T['Qga']}</b>", f"<b>{_f(group['Q_all'], 0)}</b>", "kN"],
    ]))

    # ---------------------------------------------------------- settlement
    st = res["settlement"]
    parts.append(N.h2(T["sec_settlement"]))
    parts.append(N.h3(T["sec_single_s"]))
    single = st["single"]
    parts.append(_kv_table(T, [[T["s1"], _f(1000 * single["s1"]), "mm"],
                               [T["s2"], _f(1000 * single["s2"]), "mm"],
                               [T["s3"], _f(1000 * single["s3"]), "mm"],
                               [f"<b>{T['s']}</b>", f"<b>{_f(1000 * single['s'])}</b>", "mm"]]))
    parts.append(N.h3(T["sec_group_s"]))
    raft = st["raft"]
    rows = [[T["raft_z"], _f(raft["z_raft"]), "m"], [T["raft_q"], _f(raft["q"], 1), "kPa"],
            [T["raft_c"], _f(1000 * raft["consolidation"]), "mm"],
            [T["raft_e"], _f(1000 * raft["elastic"]), "mm"],
            [T["raft_sh"], _f(1000 * raft["shortening"]), "mm"],
            [f"{L['gs_raft']} — {T['raft_total']}", _f(1000 * raft["total"]), "mm"],
            [L["gs_vesic"], _f(1000 * st["vesic"]), "mm"]]
    if st["meyerhof"]:
        rows.append([L["gs_meyerhof"], _f(1000 * st["meyerhof"]["s"]), "mm"])
    parts.append(_kv_table(T, rows))
    if raft["rows"]:
        parts.append(N.h3(T["sec_raft"]))
        rows = [[_f(r["z"]), _f(r["h"]), _esc(r["layer"]), _f(r["sigma0"], 1),
                 _f(r["dsigma"], 1), L["model_" + r["kind"]], _f(1000 * r["s"])]
                for r in raft["rows"]]
        parts.append(_table([T["z"], T["h"], T["layer"], T["s0"], T["ds"], T["kind"],
                             T["ds_mm"]], rows, [10, 8, 24, 14, 14, 18, 12]))

    # ---------------------------------------------------------- checks
    parts.append(N.h2(T["sec_checks"]))
    checks = res["checks"]
    rows = [[T["c_pile"], _f(checks["pile"]["actual"]), _f(checks["pile"]["allowable"]),
             _status(T, checks["pile"]["status"])],
            [T["c_group"], _f(checks["group"]["actual"]), _f(checks["group"]["allowable"]),
             _status(T, checks["group"]["status"])],
            [T["c_settle"], _f(checks["settlement"]["actual"], 1),
             _f(checks["settlement"]["allowable"], 1), _status(T, checks["settlement"]["status"])]]
    parts.append(_table([T["check"], T["actual"], T["allowable"], T["status"]], rows,
                        [40, 20, 20, 20]))
    length = res["required_length"]
    text = (L["res_required_length"].format(L=length, z=a.top + length) if math.isfinite(length)
            else L["res_no_length"].format(hi=a.profile.depth - a.top))
    parts.append(f"<p>{_esc(text)}</p>")

    parts.append(N.h2(T["sec_figs"], page_break=True))
    for key in PLOT_KEYS:
        if key in figures:
            parts.append(f"<p class='lead'>{_esc(L.get('fig_' + key, key))}</p>")
            parts.append(f"<p><img src='{img_src(key)}' width='640'></p>")

    if res["warnings"]:
        parts.append(N.h2(T["sec_warn"]) + "<ul>")
        parts.extend(f"<li>{_esc(warning_text(lang, w))}</li>" for w in res["warnings"])
        parts.append("</ul>")
    parts.append(N.h2(T["sec_notes"]) + "<ul class='note'>")
    parts.extend(f"<li>{_esc(n)}</li>" for n in T["notes"])
    parts.append("</ul>")
    return parts


def _socket_section(T, L, s, figures, img_src, lang, N: _Numbers, page_break: bool) -> List[str]:
    res = s.results
    parts = [N.h2(T["sec_socket"], page_break=page_break), N.h3(T["sec_socket_in"])]
    m = res["modulus"]
    hb = res["hoek_brown"]
    parts.append(_kv_table(T, [
        [T["socket_D"], _f(s.D), "m"], [T["socket_Ls"], _f(res["Ls_check"]), "m"],
        [T["socket_top"], _f(s.top), "m"], [T["rock"], _f(s.rock_depth), "m"],
        [T["overburden"], _f(res["overburden"]), "m"], [T["socket_Q"], _f(s.Q, 0), "kN"],
        [T["qu"], _f(s.qu, 1), "MPa"], [T["qu_side"], _f(res["qu_side"], 1), "MPa"],
        [T["RQD"], _f(s.RQD, 0), "%"], [T["Ei"], _f(s.Ei, 0), "MPa"],
        [T["Em"], _f(m["Em"], 0), "MPa"], [T["ratio"], _f(m["ratio"], 3), ""],
        [T["alpha_E"], _f(m["alpha_E"], 3), ""], [T["GSI"], _f(s.GSI, 0), ""],
        [T["mi"], _f(s.mi, 1), ""], [T["mb"], _f(hb["mb"], 3), ""],
        [T["s_hb"], f"{hb['s']:.2e}", ""], [T["fc"], _f(s.fc, 1), "MPa"],
        [T["Ec"], _f(s.Ec, 0), "MPa"],
    ]))
    parts.append(N.h3(T["sec_socket_side"]))
    rows = []
    for key in SOCKET_SIDE_METHODS:
        e = res["side"][key]
        rows.append([_esc(L["side_" + key] + (" *" if key == res["design"] else "")),
                     _esc(e["formula"]), _f(e["fs"], 0), _f(e["Ls_req"]), _f(e["Q_all"], 0),
                     T["yes"] if e["in_range"] else T["no"]])
    parts.append(_table([T["correlation"], T["formula"], T["fs"], T["Ls_req"], T["Q_all"],
                         T["range"]], rows, [30, 22, 11, 12, 14, 11]))
    st = res["stats"]
    parts.append(f"<p>{_esc(L['sock_stats'].format(n=st['n'], mean=st['mean'], median=st['median'], lo=st['lower'], hi=st['upper']))}</p>")
    parts.append(N.h3(T["sec_socket_base"]))
    rows = [[_esc(L["base_" + key]), _esc(res["base"][key]["formula"]),
             _f(res["base"][key]["qb"] / 1000.0)] for key in SOCKET_BASE_METHODS]
    parts.append(_table([T["method"], T["base_formula"], T["qb_mpa"]], rows, [45, 30, 25]))

    c = res["check"]
    parts.append(N.h3(T["sec_socket_design"]))
    parts.append(_kv_table(T, [
        [T["fs_design"], _f(res["fs_design"], 0), "kPa"],
        [T["qb_design"], _f(res["qb_design"] / 1000.0), "MPa"],
        [T["FS_side"], _f(s.FS_side), ""], [T["FS_base"], _f(s.FS_base), ""],
        [T["Ls_need"], _f(res["Ls_req"]), "m"], [T["Ls_min"], _f(res["Ls_min"]), "m"],
        [f"<b>{T['Ls_design']}</b>", f"<b>{_f(res['Ls_design'])}</b>", "m"],
        [T["socket_Ls"], _f(res["Ls_check"]), "m"],
        [T["Qs_l"], _f(c["Qs"], 0), "kN"], [T["Qb_l"], _f(c["Qb"], 0), "kN"],
        [T["W"], _f(c["W"], 0), "kN"], [T["Qa"].split("=")[0].strip(), _f(c["Q_all"], 0), "kN"],
        [T["status"], _status(T, c["status"]), ""],
    ]))
    se = res["settlement"]
    parts.append(N.h3(T["sec_socket_s"]))
    parts.append(_kv_table(T, [
        [T["shortening"], _f(1000 * se["shortening"]), "mm"],
        [T["rw"], _f(1000 * se["rw"]["total"]), "mm"],
        [T["base_share"], _f(100 * se["rw"]["base_share"], 1), "%"],
        [T["rw_side"], _f(1000 * se["rw_side"]["total"]), "mm"],
        [T["vesic"], _f(1000 * se["vesic"]["total"]), "mm"],
    ]))
    parts.append(N.h3(T["sec_socket_figs"]))
    for key in SOCKET_PLOT_KEYS:
        if key in figures:
            parts.append(f"<p class='lead'>{_esc(L.get('fig_' + key, key))}</p>")
            parts.append(f"<p><img src='{img_src(key)}' width='640'></p>")
    if res["warnings"]:
        parts.append(N.h3(T["sec_socket_warn"]) + "<ul>")
        parts.extend(f"<li>{_esc(warning_text(lang, w))}</li>" for w in res["warnings"])
        parts.append("</ul>")
    parts.append(N.h3(T["sec_socket_notes"]) + "<ul class='note'>")
    parts.extend(f"<li>{_esc(n)}</li>" for n in T["socket_notes"])
    parts.append("</ul>")
    return parts


def _study_section(T, L, study, figures, img_src, N: _Numbers) -> List[str]:
    parts = [N.h2(T["sec_study"], page_break=True)]
    s = study.summary
    parts.append(_kv_table(T, [
        [T["study_method"], L[f"sampling_{study.method}"], ""],
        [T["study_n"], str(s.get("n_total", 0)), ""],
        [T["study_ok"], str(s.get("n_ok", 0)), ""],
    ]))
    parts.append(N.h3(T["sec_study_vars"]))
    rows = []
    for v in study.variables:
        if v.mode == "range":
            rows.append([_esc(v.label), T["range_mode"], _f(v.vmin, 3), _f(v.vmax, 3), "—", "—",
                         "—", v.n_points])
        else:
            rows.append([_esc(v.label), T["dist"], "—", "—", L[f"dist_{v.dist}"],
                         _f(v.mean, 3), _f(v.cov, 3), "—"])
    parts.append(_table([T["variable"], T["mode"], T["min"], T["max"], T["dist"], T["mean"],
                         T["cov"], "n"], rows, [28, 12, 10, 10, 14, 10, 8, 8]))
    parts.append(N.h3(T["sec_study_stats"]))
    rows = [[_esc(L.get(f"out_{k}", k)), st["n"], _f(st["mean"], 3), _f(st["std"], 3),
             _f(st["p5"], 3), _f(st["p50"], 3), _f(st["p95"], 3)]
            for k, _ in OUTPUTS if k in s.get("stats", {}) for st in [s["stats"][k]]]
    parts.append(_table([T["output"], "n", T["mean"], T["std"], "P5", "P50", "P95"],
                        rows, [34, 8, 12, 12, 11, 11, 12]))
    if s.get("reliability"):
        parts.append(N.h3(T["sec_study_rel"]))
        rows = []
        for name, r in s["reliability"].items():
            rows.append([_esc(L.get(name, name)), r["n"], r["n_fail"], f"{r['pf']:.3g}",
                         f"[{r['pf_lo']:.2g}, {r['pf_hi']:.2g}]",
                         _esc(study_plots._beta_text(r["beta"], r["n"]))])
        parts.append(_table([T["criterion"], "n", T["n_fail"], T["pf"], T["pf_ci"],
                             T["beta_idx"]], rows, [30, 10, 12, 14, 20, 14]))
    if s.get("spearman"):
        parts.append(N.h3(T["sec_study_sens"]))
        outs = [k for k, _ in OUTPUTS if k in s["spearman"]]
        rows = [[_esc(v.label)] + [f"{s['spearman'][k].get(v.path, float('nan')):+.2f}"
                                   for k in outs] for v in study.variables]
        parts.append(_table([T["variable"]] + [L.get(f"out_{k}", k) for k in outs], rows))
    parts.append(N.h3(T["sec_study_figs"]))
    for key in figures:
        if key.startswith("study_"):
            parts.append(f"<p><img src='{img_src(key)}' width='640'></p>")
    return parts


def build_html(analysis, lang: str, figures: Dict[str, bytes], img_src=None,
               study=None, socket=None) -> str:
    """
    Builds the report HTML. img_src(key) -> value for the <img src> attribute;
    the default embeds base64 data URIs (the self-contained HTML file). The PDF
    and DOCX writers pass a `fig://<key>` mapper and resolve the keys against
    the figure dictionary themselves.

    Either the pile analysis or the socket may be missing, not both.
    """
    if analysis is None and socket is None:
        raise ValueError("nothing to report")
    T = TEXTS.get(lang, TEXTS["en"])
    L = TRANSLATIONS.get(lang, TRANSLATIONS["en"])
    source = analysis if analysis is not None else socket
    info = source.config.get("project_info", {})
    if img_src is None:
        def img_src(key):
            return "data:image/png;base64," + base64.b64encode(figures[key]).decode("ascii")

    parts = [f"<html><head><meta charset='utf-8'><style>{_CSS}</style></head><body>"]
    parts.append(f"<h1>{_esc(info.get('title') or T['title'])}</h1>")
    stamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    meta = f"{T['date']}: {stamp}"
    if info.get("analyst"):
        meta += f" &nbsp;·&nbsp; {T['analyst']}: {_esc(info['analyst'])}"
    meta += f" &nbsp;·&nbsp; {T['software']}: {APP_NAME} v{APP_VERSION}"
    parts.append(f"<p class='meta'>{meta}</p>")

    N = _Numbers()
    if analysis is not None:
        parts.extend(_pile_section(T, L, analysis, figures, img_src, lang, N))
    if socket is not None:
        parts.extend(_socket_section(T, L, socket, figures, img_src, lang, N,
                                     page_break=analysis is not None))
    if study is not None and study.rows:
        parts.extend(_study_section(T, L, study, figures, img_src, N))
    parts.append("</body></html>")
    return "\n".join(parts)


# ----------------------------------------------------------------------
# exporters
# ----------------------------------------------------------------------
def _all_figures(analysis, lang, study, socket):
    figures = render_figures(analysis, lang)
    figures.update(render_socket_figures(socket, lang))
    figures.update(render_study_figures(study, lang))
    return figures


def export_html(path: str, analysis, lang: str, study=None, socket=None) -> None:
    figures = _all_figures(analysis, lang, study, socket)
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(build_html(analysis, lang, figures, study=study, socket=socket))


def export_pdf(path: str, analysis, lang: str, study=None, socket=None) -> None:
    """PDF through reportlab; the figures travel as `fig://<key>` references."""
    from . import pdf as pdf_writer

    T = TEXTS.get(lang, TEXTS["en"])
    source = analysis if analysis is not None else socket
    info = source.config.get("project_info", {})
    figures = _all_figures(analysis, lang, study, socket)
    html_text = build_html(analysis, lang, figures, img_src=lambda k: f"fig://{k}",
                           study=study, socket=socket)
    pdf_writer.html_to_pdf(path, html_text, figures,
                           title=info.get("title", T["title"]),
                           author=info.get("analyst", ""),
                           footer=f"{info.get('title', '')} — {APP_NAME} v{APP_VERSION}")


def export_docx(path: str, analysis, lang: str, study=None, socket=None) -> None:
    """Word report; requires python-docx."""
    try:
        import docx
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:
        raise RuntimeError("python-docx is not installed (pip install python-docx).") from exc
    import re

    figures = _all_figures(analysis, lang, study, socket)
    html_text = build_html(analysis, lang, figures, img_src=lambda k: f"fig://{k}",
                           study=study, socket=socket)
    d = docx.Document()
    d.styles["Normal"].font.size = Pt(10)
    d.styles["Normal"].font.color.rgb = RGBColor(0x14, 0x14, 0x13)
    # the interface's look: serif headings, the title in ink, sections in terracotta
    for name, colour in (("Title", (0x14, 0x14, 0x13)), ("Heading 1", (0xC6, 0x61, 0x3F)),
                         ("Heading 2", (0x14, 0x14, 0x13))):
        style = d.styles[name]
        style.font.name = "Georgia"
        style.font.bold = False
        style.font.color.rgb = RGBColor(*colour)

    tokens = re.split(r"(<h1>.*?</h1>|<h2[^>]*>.*?</h2>|<h3>.*?</h3>|<table[^>]*>.*?</table>|"
                      r"<p[^>]*>.*?</p>|<li>.*?</li>)", html_text, flags=re.S)

    def strip(s):
        return html.unescape(re.sub(r"<[^>]+>", "", re.sub(r"<br\s*/?>", "\n", s))).strip()

    for tok in tokens:
        if tok.startswith("<h1>"):
            d.add_heading(strip(tok), level=0)
        elif tok.startswith("<h2"):
            d.add_heading(strip(tok), level=1)
        elif tok.startswith("<h3>"):
            d.add_heading(strip(tok), level=2)
        elif tok.startswith("<li>"):
            d.add_paragraph(strip(tok), style="List Bullet")
        elif tok.startswith("<table"):
            rows = re.findall(r"<tr>(.*?)</tr>", tok, flags=re.S)
            cells = [re.findall(r"<t[hd][^>]*>(.*?)</t[hd]>", r, flags=re.S) for r in rows]
            if cells:
                table = d.add_table(rows=len(cells), cols=len(cells[0]))
                table.style = "Light Grid Accent 1"
                for i, row in enumerate(cells):
                    for j, c in enumerate(row):
                        if j < len(table.columns):
                            table.cell(i, j).text = strip(c)
        elif tok.startswith("<p"):
            m = re.search(r"src='fig://([a-z_]+)'", tok)
            if m:
                d.add_picture(io.BytesIO(figures[m.group(1)]), width=Inches(6.3))
            else:
                text = strip(tok)
                if text:
                    d.add_paragraph(text)
    d.save(path)


def export_report(path: str, analysis, lang: str, study=None, socket=None) -> None:
    ext = path.lower().rsplit(".", 1)[-1]
    if ext == "pdf":
        export_pdf(path, analysis, lang, study, socket)
    elif ext == "docx":
        export_docx(path, analysis, lang, study, socket)
    else:
        export_html(path if ext in ("html", "htm") else path + ".html", analysis, lang, study,
                    socket)
